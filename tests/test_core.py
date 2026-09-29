from io import BytesIO
import copy
import json
import time
import uuid
from concurrent.futures import ThreadPoolExecutor
from dataclasses import replace
import httpx
import pytest
from docx import Document as WordDocument
from pypdf import PdfReader
from app import providers
from app.db import reserve, finish_job, expire_jobs, usage, put_state, Session, Job
from app.parsing import parse_upload
from app.exports import export_pdf, export_docx
from app.scoring import assess
from app.sample import sample_document

JD = 'Java Spring Boot React Docker PostgreSQL REST APIs and automated testing.'

def test_concurrent_daily_reservations_are_bounded():
    def attempt(n):
        try: return reserve(uuid.uuid4().hex, str(n))[0].id
        except ValueError: return None
    with ThreadPoolExecutor(max_workers=10) as workers:
        ids = [x for x in workers.map(attempt, range(30)) if x]
    assert len(ids) == 15
    for ident in ids:
        finish_job(ident, True); finish_job(ident, True)
    assert usage()['used'] == 15 and usage()['reserved'] == 0

def test_interrupted_jobs_release_reservation():
    job, _ = reserve(uuid.uuid4().hex, 'resume')
    with Session.begin() as s: s.get(Job, job.id).created = time.time() - 700
    expire_jobs()
    assert usage()['reserved'] == 0 and usage()['used'] == 0

def test_docx_import_keeps_line_breaks_images_and_skill_labels():
    from PIL import Image
    from docx.shared import Inches
    doc = WordDocument()
    doc.add_paragraph('Alex Example\nJava Developer\nEmail: example@example.com')
    doc.add_paragraph('TECHNICAL SKILLS')
    row = doc.add_table(rows=1, cols=2).rows[0]
    row.cells[0].text = 'Languages'; row.cells[1].text = 'Java, Python'
    img = BytesIO(); Image.new('RGB', (100,100), 'white').save(img, format='PNG'); img.seek(0)
    doc.add_picture(img, width=Inches(0.5))
    raw = BytesIO(); doc.save(raw)
    parsed, images = parse_upload('resume.docx', raw.getvalue())
    assert [b.kind for b in parsed.blocks[:3]] == ['name','subtitle','contact']
    assert any(b.kind == 'skill' and 'Languages:' in b.text for b in parsed.blocks)
    assert len(images) == 1

def test_prompt_compact_contact_exclusion_and_numbers_guard(document):
    messages, evidence = providers.evidence_prompt(document, JD)
    assert 'alex@example.com' not in messages[1]['content']
    assert len(evidence) <= 24
    ident = evidence[0]['id']
    result = providers.validate_result(json.dumps({'priority_ids':[ident,'invented'], 'suggestions':[{'block_id':ident,'proposed':'Improved throughput by 99%'}]}), evidence)
    assert result['priority_ids'] == [ident] and result['suggestions'] == []
    with pytest.raises(ValueError): providers.validate_result('[]', evidence)

def mock_providers(monkeypatch, handler):
    cfg = replace(providers.config, groq_key='test-groq', groq_free=True, openrouter_key='test-router', ollama_url='')
    monkeypatch.setattr(providers, 'config', cfg)
    original = httpx.Client
    monkeypatch.setattr(providers.httpx, 'Client', lambda **kw: original(transport=httpx.MockTransport(handler), **kw))

def test_quota_fallback_skips_other_models_in_exhausted_account(monkeypatch, document):
    calls = []
    def handle(req):
        calls.append(str(req.url))
        if req.url.host == 'api.groq.com':
            return httpx.Response(429, headers={'retry-after':'60'}, json={'error':'rate limit'})
        if req.url.path.endswith('/models'):
            return httpx.Response(200, json={'data':[{'id':'openai/gpt-oss-20b:free','pricing':{'prompt':'0','completion':'0'}}]})
        body=json.loads(req.content)
        assert body['model'].endswith(':free') and body['provider']['max_price']['completion'] == 0
        evidence=json.loads(body['messages'][1]['content'])['evidence_excerpts']
        return httpx.Response(200, json={'choices':[{'finish_reason':'stop','message':{'content':json.dumps({'priority_ids':[evidence[0]['id']],'suggestions':[]})}}]})
    mock_providers(monkeypatch, handle)
    result = providers.run_ai(document, JD, lambda *a: None)
    assert result['provider'] == 'OpenRouter'
    assert len([x for x in calls if 'api.groq.com' in x]) == 1
    assert providers.public_status()[0]['status'] == 'cooling_down'

def test_paid_catalog_models_never_called(monkeypatch, document):
    calls = []
    put_state('settings', {'fallback':True, 'enabled':['openrouter']})
    def handle(req):
        calls.append(req.url.path)
        return httpx.Response(200, json={'data':[{'id':'openai/gpt-oss-20b:free','pricing':{'prompt':'0.01','completion':'0.01'}}]})
    mock_providers(monkeypatch, handle)
    with pytest.raises(providers.ProviderFailure): providers.run_ai(document, JD, lambda *a:None)
    assert calls == ['/api/v1/models']

def test_disabled_fallback_never_uses_secondary_during_cooldown(monkeypatch, document):
    put_state('settings', {'fallback':False,'enabled':['groq','openrouter']})
    put_state('provider_status', {'groq':{'retry_at': time.time()+100, 'status':'cooling_down'}})
    def unexpected(req): raise AssertionError('No remote call should occur')
    mock_providers(monkeypatch, unexpected)
    with pytest.raises(providers.ProviderFailure): providers.run_ai(document, JD, lambda *a:None)

def test_unrelated_job_is_not_forced_above_ninety(document):
    result = assess(document, 'A surgeon experienced in cardiology, neurosurgery, transplant surgery and clinical trials is required.')
    assert result['score'] < 90

@pytest.mark.parametrize('target',[5,6,7])
def test_long_resume_pdf_target_and_editable_docx(target):
    from app.tailoring import tailor
    doc = sample_document(); doc['target_pages'] = target
    tailored, _ = tailor(doc, JD)
    data, plan = export_pdf(tailored)
    pdf = PdfReader(BytesIO(data))
    assert len(pdf.pages) == len(plan.pages) and len(pdf.pages) <= target
    assert len(pdf.pages) >= 5
    assert all(len(p.extract_text()) > 100 for p in pdf.pages)
    word, _ = export_docx(tailored)
    loaded=WordDocument(BytesIO(word))
    assert 'Alex Morgan' in loaded.paragraphs[0].text
    assert all(not p._p.xpath('.//w:br[@w:type="page"]') for p in loaded.paragraphs)
