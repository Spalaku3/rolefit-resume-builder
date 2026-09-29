"""Measured PDF/preview and naturally flowing editable DOCX; never add filler."""
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from xml.sax.saxutils import escape
import re
import os
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import Paragraph, Table, TableStyle, Image, HRFlowable, KeepTogether
from reportlab.lib.pagesizes import letter, A4
from docx import Document as WordDocument
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from .schemas import Document

ACCENT = '#9F2435'
MARGIN = 36


def _font_pair(family='Aptos'):
    serif = family == 'Times New Roman'
    candidates = [(os.getenv('RESUME_FONT_REGULAR', ''), os.getenv('RESUME_FONT_BOLD', ''))]
    if family in {'Aptos', 'Calibri'}:
        candidates += [('/usr/share/fonts/truetype/crosextra/Carlito-Regular.ttf', '/usr/share/fonts/truetype/crosextra/Carlito-Bold.ttf'),
                       ('C:/Windows/Fonts/calibri.ttf', 'C:/Windows/Fonts/calibrib.ttf')]
    face = 'Serif' if serif else 'Sans'
    for folder in ['liberation2', 'liberation']:
        candidates.append((f'/usr/share/fonts/truetype/{folder}/Liberation{face}-Regular.ttf', f'/usr/share/fonts/truetype/{folder}/Liberation{face}-Bold.ttf'))
    candidates += [('C:/Windows/Fonts/times.ttf' if serif else 'C:/Windows/Fonts/arial.ttf',
                    'C:/Windows/Fonts/timesbd.ttf' if serif else 'C:/Windows/Fonts/arialbd.ttf')]
    for regular, bold in candidates:
        if regular and Path(regular).is_file() and Path(bold).is_file():
            import hashlib
            name = 'RoleFit' + hashlib.sha256(regular.encode()).hexdigest()[:8]
            bold_name = name + 'Bold'
            if name not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont(name, regular))
                pdfmetrics.registerFont(TTFont(bold_name, bold))
                pdfmetrics.registerFontFamily(name, normal=name, bold=bold_name, italic=name, boldItalic=bold_name)
            return name, bold_name
    return ('Times-Roman', 'Times-Bold') if serif else ('Helvetica', 'Helvetica-Bold')


def markup(block):
    text = escape(block['text'])
    # Escape first; never interpret uploaded text as HTML/ReportLab markup.
    for phrase in sorted(set(block.get('bold', [])), key=len, reverse=True)[:30]:
        quoted = escape(phrase)
        if quoted and quoted in text and '<b>' not in quoted:
            # Avoid replacing inside markup from an earlier larger phrase.
            parts = re.split(r'(<b>.*?</b>)', text)
            text = ''.join(p if p.startswith('<b>') else p.replace(quoted, '<b>' + quoted + '</b>') for p in parts)
    return text.replace('\n', '<br/>')


@dataclass
class Item:
    block: dict
    flow: object
    height: float
    before: float
    after: float

@dataclass
class Plan:
    pages: list
    font_size: float
    leading: float
    page_size: tuple
    width: float
    header_height: float
    header: object
    header_blocks: list
    images: list
    warnings: list


def _header(blocks, images, width, regular, bold, size, template):
    header_blocks = [b for b in blocks if b['kind'] in {'name', 'subtitle', 'contact'}]
    header_images = [a for a in images if a.get('placement', 'header') == 'header'][:3] if template != 'ats' else []
    left = []
    for b in header_blocks:
        st = ParagraphStyle('header', fontName=bold if b['kind'] in {'name', 'subtitle'} else regular,
                            fontSize=size + 1 if b['kind'] == 'name' else size,
                            leading=(size + 1) * 1.12, spaceAfter=1)
        left.append(Paragraph(markup(b), st))
    if not left:
        left = [Paragraph('Resume', ParagraphStyle('header', fontName=bold, fontSize=size + 2, leading=size + 5))]
    if header_images:
        pictures = []
        for a in header_images:
            im = Image(BytesIO(a['data']))
            scale = min(64 / im.imageWidth, 66 / im.imageHeight)
            im.drawWidth = im.imageWidth * scale; im.drawHeight = im.imageHeight * scale
            pictures.append(im)
        right = Table([pictures], colWidths=[70] * len(pictures))
        right.setStyle(TableStyle([('VALIGN', (0, 0), (-1, -1), 'TOP'), ('LEFTPADDING', (0, 0), (-1, -1), 2), ('RIGHTPADDING', (0, 0), (-1, -1), 2)]))
        right_width = 70 * len(pictures)
        table = Table([[left, right]], colWidths=[width - right_width, right_width])
    else:
        table = Table([[left]], colWidths=[width])
    table.setStyle(TableStyle([('VALIGN', (0, 0), (-1, -1), 'TOP'), ('LEFTPADDING', (0, 0), (-1, -1), 0),
                               ('RIGHTPADDING', (0, 0), (-1, -1), 0), ('TOPPADDING', (0, 0), (-1, -1), 0), ('BOTTOMPADDING', (0, 0), (-1, -1), 3)]))
    _, h = table.wrap(width, 900)
    return table, h + 13, header_blocks


def _items(document, images, size, width):
    regular, bold = _font_pair(document.get('font', 'Aptos'))
    leading = size * 1.16
    template = document.get('template', 'reference')
    items = []
    for original in document['blocks']:
        if original['kind'] in {'name', 'subtitle', 'contact'}:
            continue
        # Bound individual paragraph heights while retaining every word.
        words = original['text'].split()
        chunks = []
        while words:
            chunk, words = words[:115], words[115:]
            chunks.append(' '.join(chunk))
        for index, text in enumerate(chunks or [original['text']]):
            b = {**original, 'text': text}
            k = b['kind']
            is_heading = k in {'heading', 'subheading'}
            st = ParagraphStyle('p', fontName=bold if is_heading else regular, fontSize=size,
                                leading=leading, alignment=TA_LEFT,
                                textColor=colors.HexColor(ACCENT) if k == 'heading' and template == 'reference' else colors.black,
                                leftIndent=12 if k == 'bullet' else 0,
                                firstLineIndent=-9 if k == 'bullet' else 0)
            before = 11 if k == 'heading' else 6 if k == 'subheading' else 0
            after = 7 if k == 'heading' else 4 if k == 'subheading' else 3
            txt = markup(b)
            if k == 'heading':
                txt = '<u>' + txt.upper() + '</u>' if template == 'reference' else txt.upper()
            if k == 'bullet':
                txt = '\u2022 ' + txt
            if k == 'skill' and template == 'reference' and ':' in b['text']:
                label, content = b['text'].split(':', 1)
                p1 = Paragraph('<b>' + escape(label) + '</b>', st)
                p2 = Paragraph(escape(content.strip()), st)
                flow = Table([[p1, p2]], colWidths=[width * 0.28, width * 0.72])
                flow.setStyle(TableStyle([('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#555555')),
                                         ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'), ('LEFTPADDING', (0, 0), (-1, -1), 5),
                                         ('RIGHTPADDING', (0, 0), (-1, -1), 5), ('TOPPADDING', (0, 0), (-1, -1), 3),
                                         ('BOTTOMPADDING', (0, 0), (-1, -1), 3)]))
                after = 0
            else:
                flow = Paragraph(txt, st)
            _, h = flow.wrap(width, 2000)
            items.append(Item(b, flow, h, before, after))
    if template != 'ats':
        for asset in [a for a in images if a.get('placement') == 'end']:
            im = Image(BytesIO(asset['data']))
            factor = min(220 / im.imageWidth, 150 / im.imageHeight)
            im.drawWidth = im.imageWidth * factor; im.drawHeight = im.imageHeight * factor
            items.append(Item({'kind': 'image', 'id': asset['id'], 'text': asset['label']}, im, im.drawHeight, 10, 6))
    return items, leading


def make_plan(document, images=None):
    Document.model_validate(document)
    images = images or []
    size_page = A4 if document.get('paper') == 'a4' else letter
    width = size_page[0] - 2 * MARGIN
    regular, bold = _font_pair(document.get('font', 'Aptos'))
    target = document.get('target_pages', 6)
    best = None
    for size in [12, 11.5, 11, 10.5]:
        header, hh, hb = _header(document['blocks'], images, width, regular, bold, size, document.get('template'))
        items, leading = _items(document, images, size, width)
        usable = size_page[1] - MARGIN - 40
        pages, page, used = [], [], hh
        for i, item in enumerate(items):
            need = item.before + item.height + item.after
            # Keep each heading with the next substantive line; do not strand it.
            reserve = 0
            if item.block['kind'] in {'heading', 'subheading'} and i + 1 < len(items):
                nxt = items[i + 1]
                reserve = nxt.before + nxt.height + nxt.after
                if nxt.block['kind'] == 'subheading' and i + 2 < len(items):
                    reserve += items[i + 2].before + items[i + 2].height
            if used + need + reserve > usable and (page or not pages):
                pages.append(page); page = []; used = 0
            page.append(item); used += need
        if page or not pages:
            pages.append(page)
        warnings = []
        if len(pages) < target:
            warnings.append(f'This content uses {len(pages)} pages at a readable size. No filler or fabricated experience was added to reach {target}.')
        if len(pages) > target:
            warnings.append(f'Current content needs {len(pages)} pages. Edit or shorten it to reach your {target}-page target; nothing is silently cut during export.')
        if document.get('font') == 'Aptos':
            warnings.append('DOCX requests Aptos. PDF uses an installed sans-serif fallback unless custom font paths are configured. Font files are not bundled.')
        best = Plan(pages, size, leading, size_page, width, hh, header, hb, images, warnings)
        if len(pages) <= target:
            break
    return best


def plan_info(document, images=None):
    p = make_plan(document, images)
    return {'pages': len(p.pages), 'target_pages': document.get('target_pages', 6), 'font_size': p.font_size,
            'warnings': p.warnings, 'pdf_font': 'Installed serif fallback' if document.get('font') == 'Times New Roman' else 'Installed sans-serif fallback',
            'within_target': 5 <= len(p.pages) <= 7}


def export_pdf(document, images=None):
    plan = make_plan(document, images)
    out = BytesIO()
    c = canvas.Canvas(out, pagesize=plan.page_size, pageCompression=1)
    c.setTitle('Resume')
    c.setAuthor('')
    regular, bold = _font_pair(document.get('font', 'Aptos'))
    for n, items in enumerate(plan.pages):
        y = plan.page_size[1] - MARGIN
        if n == 0:
            plan.header.drawOn(c, MARGIN, y - plan.header_height + 13)
            y -= plan.header_height
            if document.get('template') == 'reference':
                c.setStrokeColor(colors.HexColor('#79545A')); c.setLineWidth(1.7)
                c.line(MARGIN, y + 3, MARGIN + plan.width, y + 3)
                c.setLineWidth(0.7); c.line(MARGIN, y, MARGIN + plan.width, y)
        for item in items:
            y -= item.before
            if item.block['kind'] == 'heading' and document.get('template') == 'reference':
                c.setStrokeColor(colors.HexColor('#79545A')); c.setLineWidth(0.6)
                c.line(MARGIN, y + 4, MARGIN + plan.width, y + 4)
            item.flow.drawOn(c, MARGIN, y - item.height)
            y -= item.height + item.after
        c.setFont(regular, 8)
        c.setFillColor(colors.HexColor('#555555'))
        c.drawString(MARGIN, 20, f'Page {n + 1} of {len(plan.pages)}')
        c.showPage()
    c.save()
    return out.getvalue(), plan


def _border(paragraph, color='79545A'):
    pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement('w:pBdr')
    b = OxmlElement('w:bottom')
    for k, v in {'val': 'double', 'sz': '5', 'space': '3', 'color': color}.items():
        b.set(qn('w:' + k), v)
    borders.append(b); pr.append(borders)


def _runs(paragraph, b, force_bold=False):
    text = b['text']
    phrases = [x for x in b.get('bold', []) if x in text and len(x) > 2]
    if not phrases:
        paragraph.add_run(text).bold = force_bold
        return
    regex = re.compile('(' + '|'.join(re.escape(x) for x in sorted(set(phrases), key=len, reverse=True)) + ')')
    for piece in regex.split(text):
        r = paragraph.add_run(piece)
        r.bold = force_bold or piece in phrases


def _configure_paragraph(p, size, leading, before=0, after=3):
    pf = p.paragraph_format
    pf.space_before = Pt(before); pf.space_after = Pt(after); pf.line_spacing = Pt(leading)
    pf.widow_control = False
    pf.keep_with_next = False
    pf.keep_together = False


def export_docx(document, images=None):
    plan = make_plan(document, images)
    doc = WordDocument()
    section = doc.sections[0]
    section.page_width = Pt(plan.page_size[0]); section.page_height = Pt(plan.page_size[1])
    section.top_margin = Pt(MARGIN); section.bottom_margin = Pt(32)
    section.left_margin = Pt(MARGIN); section.right_margin = Pt(MARGIN)
    section.footer_distance = Pt(12)
    style = doc.styles['Normal']
    style.font.name = document.get('font', 'Aptos')
    style.font.size = Pt(plan.font_size)
    style.paragraph_format.line_spacing = Pt(plan.leading)
    style.paragraph_format.space_after = Pt(3)
    # Do not leave theme font assignments overriding the explicit requested family.
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    for key in ('asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme'):
        fonts.attrib.pop(qn('w:' + key), None)
    fonts.set(qn('w:ascii'), document.get('font', 'Aptos'))
    fonts.set(qn('w:hAnsi'), document.get('font', 'Aptos'))
    fonts.set(qn('w:eastAsia'), document.get('font', 'Aptos'))
    fonts.set(qn('w:cs'), document.get('font', 'Aptos'))
    # OOXML font metadata helps office applications select a sensible substitute.
    from lxml import etree
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    font_part = doc.part.part_related_by(RT.FONT_TABLE)
    root = etree.fromstring(font_part.blob)
    chosen = document.get('font', 'Aptos')
    if not root.xpath('./w:font[@w:name=$family]', namespaces={'w': qn('w:font').split('}')[0][1:]}, family=chosen):
        entry = OxmlElement('w:font'); entry.set(qn('w:name'), chosen)
        family = OxmlElement('w:family'); family.set(qn('w:val'), 'roman' if chosen == 'Times New Roman' else 'swiss'); entry.append(family)
        alternate = OxmlElement('w:altName'); alternate.set(qn('w:val'), 'Liberation Serif' if chosen == 'Times New Roman' else 'Carlito'); entry.append(alternate)
        root.append(entry)
    font_part._blob = etree.tostring(root, xml_declaration=True, encoding='UTF-8', standalone=True)

    footer = section.footer.paragraphs[0]
    _configure_paragraph(footer, 8, 9, after=0)
    footer.add_run('Page ').font.size = Pt(8)
    for field_name, following in [('PAGE', ' of '), ('NUMPAGES', '')]:
        run = footer.add_run(); field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), field_name)
        run._r.addnext(field)
        if following: footer.add_run(following)
    for r in footer.runs:
        r.font.size = Pt(8)
    badge_images = [a for a in plan.images if a.get('placement', 'header') == 'header'][:3] if document.get('template') != 'ats' else []
    if badge_images:
        table = doc.add_table(rows=1, cols=1 + len(badge_images))
        table.autofit = False
        widths = [plan.width - 70 * len(badge_images)] + [70] * len(badge_images)
        for i, width in enumerate(widths):
            table.columns[i].width = Pt(width)
            table.cell(0, i).width = Pt(width)
        for cell in table.rows[0].cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            tcpr = cell._tc.get_or_add_tcPr(); margins = OxmlElement('w:tcMar')
            for side in ('top', 'left', 'bottom', 'right'):
                el = OxmlElement('w:' + side); el.set(qn('w:w'), '0'); el.set(qn('w:type'), 'dxa'); margins.append(el)
            tcpr.append(margins)
        parent = table.cell(0, 0)
        for i, block in enumerate(plan.header_blocks):
            p = parent.paragraphs[0] if i == 0 else parent.add_paragraph()
            _configure_paragraph(p, plan.font_size, (plan.font_size + 1) * 1.12, after=1)
            _runs(p, block, block['kind'] in {'name', 'subtitle'})
        from PIL import Image as PILImage
        for i, asset in enumerate(badge_images):
            p = table.cell(0, i + 1).paragraphs[0]
            _configure_paragraph(p, plan.font_size, plan.leading, after=0)
            p.paragraph_format.line_spacing = 1.0
            with PILImage.open(BytesIO(asset['data'])) as im:
                ratio = min(64 / im.width, 66 / im.height)
                p.add_run().add_picture(BytesIO(asset['data']), width=Pt(im.width * ratio), height=Pt(im.height * ratio))
        divider = doc.add_paragraph()
        _configure_paragraph(divider, 1, 1, after=6)
    else:
        for b in plan.header_blocks:
            p = doc.add_paragraph()
            _configure_paragraph(p, plan.font_size, (plan.font_size + 1) * 1.12, after=1)
            _runs(p, b, b['kind'] in {'name', 'subtitle'})
        divider = doc.add_paragraph()
        _configure_paragraph(divider, 1, 1, after=7)
    if document.get('template') == 'reference':
        _border(divider)
    asset_map = {a['id']: a for a in plan.images}
    for n, items in enumerate(plan.pages):
        # Word uses natural flow, not PDF page-break hints: different font metrics
        # otherwise create almost-empty overflow pages. The PDF is page-count exact.
        for item in items:
            b = item.block; k = b['kind']
            if k == 'image':
                a = asset_map[b['id']]
                p = doc.add_paragraph(); _configure_paragraph(p, plan.font_size, plan.leading, before=10, after=6)
                p.paragraph_format.line_spacing = 1.0
                p.add_run().add_picture(BytesIO(a['data']), width=Pt(item.flow.drawWidth), height=Pt(item.flow.drawHeight))
                continue
            if k == 'skill' and document.get('template') == 'reference' and ':' in b['text']:
                label, value = b['text'].split(':', 1)
                table = doc.add_table(rows=1, cols=2)
                table.autofit = False; table.style = 'Table Grid'
                table.columns[0].width = Pt(plan.width * 0.28); table.columns[1].width = Pt(plan.width * 0.72)
                for i, text in enumerate([label, value.strip()]):
                    cell = table.cell(0, i); cell.width = Pt(plan.width * (0.28 if i == 0 else 0.72))
                    p = cell.paragraphs[0]
                    _configure_paragraph(p, plan.font_size, plan.leading, before=1, after=1)
                    p.add_run(text).bold = i == 0
                continue
            p = doc.add_paragraph()
            _configure_paragraph(p, plan.font_size, plan.leading, item.before, item.after)
            if k in {'heading', 'subheading'}:
                p.paragraph_format.keep_with_next = True
            if k == 'bullet':
                p.paragraph_format.left_indent = Pt(12); p.paragraph_format.first_line_indent = Pt(-9)
                p.add_run('\u2022 ')
            _runs(p, {**b, 'text': b['text'].upper() if k == 'heading' else b['text']}, k in {'heading', 'subheading'})
            if k == 'heading' and document.get('template') == 'reference':
                for run in p.runs:
                    run.font.color.rgb = RGBColor.from_string(ACCENT[1:]); run.underline = True
    for paragraph in doc.element.xpath('.//w:p'):
        for element in paragraph.xpath('./w:r'):
            rp = element.get_or_add_rPr()
            rf = rp.get_or_add_rFonts()
            for attr in ('asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme'):
                rf.attrib.pop(qn('w:' + attr), None)
            for attr in ('ascii', 'hAnsi', 'cs', 'eastAsia'):
                rf.set(qn('w:' + attr), document.get('font', 'Aptos'))
    doc.core_properties.author = ''
    doc.core_properties.title = 'Resume'
    out = BytesIO(); doc.save(out)
    return out.getvalue(), plan


def export_text(document):
    return '\n\n'.join(('\u2022 ' if b['kind'] == 'bullet' else '') + b['text'] for b in document['blocks']) + '\n'
